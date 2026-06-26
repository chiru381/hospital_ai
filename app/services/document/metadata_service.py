import os
import fitz
import pandas as pd
from docx import Document
from pptx import Presentation


def get_metadata(file_path):

    ext = os.path.splitext(file_path)[1].lower()

    metadata = {
        "filename": os.path.basename(file_path),
        "extension": ext,
        "size": os.path.getsize(file_path)
    }

    if ext == ".pdf":

        pdf = fitz.open(file_path)

        metadata["pages"] = len(pdf)
        metadata["author"] = pdf.metadata.get("author")
        metadata["title"] = pdf.metadata.get("title")

        pdf.close()

    elif ext == ".docx":

        doc = Document(file_path)

        metadata["paragraphs"] = len(doc.paragraphs)

    elif ext == ".xlsx":

        xls = pd.ExcelFile(file_path)

        metadata["sheets"] = xls.sheet_names

    elif ext == ".pptx":

        prs = Presentation(file_path)

        metadata["slides"] = len(prs.slides)

    return metadata